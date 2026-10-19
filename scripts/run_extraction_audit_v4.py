#!/usr/bin/env python3
"""Extraction audit v0.4: the corrected step derivation checked by a person, on difficult constructs and on windows
of the longest proofs.

  python scripts/run_extraction_audit_v4.py --dev
  python scripts/run_extraction_audit_v4.py --register
  python scripts/run_extraction_audit_v4.py --import-hand ANSWERS.xlsx
  python scripts/run_extraction_audit_v4.py --compare
  python scripts/run_extraction_audit_v4.py --check-committed

extraction-audit-v0.1 to v0.3 compared derived graphs with reconstructions by language-model agents. Two reviews
asked for a short audit by a person, of difficult syntax and of very long proofs. This audit draws, among proofs
that no earlier audit sampled, two proofs of 5 to 30 steps for each of six constructs (`simpa ... using`, a `by`
nested in a term, case analysis by `induction`, `cases`, `rcases` or `obtain`, `calc`, `<;>` or `all_goals`, and
`try`, `first` or `repeat` written in the proof), and a window of 25 consecutive steps in each of three proofs of
more than 150 steps. A person rebuilds each graph, or the window's part of it, from the source and Lean's goal
display, following the bilingual kit in `experiments/extraction-audit-v0.4/human-kit`. The comparison is
extraction-audit-v0.1's, restricted to the window's lines for a window, where it also compares which of the window's
steps consume a goal created outside it.
"""

from __future__ import annotations

import argparse
import collections
import json
import random
import re
import sys
import time
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))
import count_linearizations_v2 as counter  # noqa: E402
import run_extraction_audit as v1  # noqa: E402
import run_extraction_audit_v2 as v2  # noqa: E402
import run_extraction_audit_v3 as v3  # noqa: E402
from run_linearizations import sha256_file, write_lf  # noqa: E402
from run_search import generated  # noqa: E402

ROOT = v1.ROOT
EXPERIMENT = ROOT / "experiments" / "extraction-audit-v0.4"
PREREG = EXPERIMENT / "preregistration.json"
SAMPLE = EXPERIMENT / "sample.json"
KIT = EXPERIMENT / "human-kit"
GUIDE = KIT / "README.md"
INSTRUCTIONS = KIT / "instructions.md"
WORKSHEET = KIT / "worksheet.md"
ANSWERS = KIT / "answers.xlsx"
RECONSTRUCTIONS = EXPERIMENT / "reconstructions-by-hand.json"
ADJUDICATION = EXPERIMENT / "adjudication.json"
RESULTS = EXPERIMENT / "results.jsonl"
SUMMARY = EXPERIMENT / "summary.json"
IMPLEMENTATIONS = {"counter": ROOT / "scripts" / "count_linearizations_v2.py",
                   "base": ROOT / "scripts" / "count_linearizations.py",
                   "v1": ROOT / "scripts" / "run_extraction_audit.py",
                   "v3": ROOT / "scripts" / "run_extraction_audit_v3.py", "runner": Path(__file__).resolve()}
SEED = 20261018
STEPS = (5, 30)
PER_CONSTRUCT = 2
LONG_MIN = 151
WINDOWS = 3
WINDOW = 25
CONTEXT_LINES = 15
ROWS = 45  # numbered rows of the Steps and Edges tables in each answer sheet
CORPORA = v3.CORPORA
TRY_WORDS = re.compile(r"(?<![\w.'])(try|first|repeat'?)(?![\w'])")
CONSTRUCTS = {
    "simpa-using": "a `simpa` node whose line, or the next, has `using`",
    "nested-by": "a `by` node nested in a term (a `Lean.Parser.Term.byTactic` node that is not the root)",
    "case-analysis": "an `induction`, `cases`, `rcases` or `obtain` node",
    "calc": "a `calc` tactic node (`Lean.calcTactic`)",
    "all-goals": "a `<;>`, `all_goals` or `any_goals` node",
    "try-first-repeat": "`try`, `first`, `repeat` or `repeat'` written in the proof's source, comments removed",
}
CASE_KINDS = {f"Lean.Parser.Tactic.{k}" for k in ("induction", "cases", "rcases", "obtain")}
ALL_GOALS_KINDS = {"Lean.Parser.Tactic.«tactic_<;>_»", "Lean.Parser.Tactic.allGoals", "Lean.Parser.Tactic.anyGoals"}


# ---------------------------------------------------------------- the sample


def audited() -> set[tuple[str, str]]:
    """Every proof an earlier extraction audit sampled."""
    return v3.audited() | {(e["module"], e["declaration"]) for e in json.loads(v3.SAMPLE.read_text(encoding="utf-8"))}


def proof_lines(entry: dict[str, Any]) -> list[str]:
    lines = v2.source_lines(entry["source"])
    return list(lines[entry["byLine"] - 1:entry["lastLine"]])


def constructs_of(entry: dict[str, Any], record: dict[str, Any]) -> set[str]:
    kinds = [(n["kind"], n["line"], n["parent"]) for n in record["nodes"]]
    lines = v2.source_lines(entry["source"])
    found = set()
    for kind, line, _ in kinds:
        if kind == "Lean.Parser.Tactic.simpa" and line is not None:
            if any("using" in lines[k - 1] for k in (line, line + 1) if 0 < k <= len(lines)):
                found.add("simpa-using")
    if any(kind == "Lean.Parser.Term.byTactic" and parent is not None for kind, _, parent in kinds):
        found.add("nested-by")
    if any(kind in CASE_KINDS for kind, _, _ in kinds):
        found.add("case-analysis")
    if any(kind == "Lean.calcTactic" for kind, _, _ in kinds):
        found.add("calc")
    if any(kind in ALL_GOALS_KINDS for kind, _, _ in kinds):
        found.add("all-goals")
    if TRY_WORDS.search(v2.COMMENT.sub("", "\n".join(proof_lines(entry)))):
        found.add("try-first-repeat")
    return found


def pool() -> list[dict[str, Any]]:
    """Proofs with one root `by` block and a name that is not generated, outside every earlier audit sample."""
    done = audited()
    out = []
    for corpus, paths in CORPORA.items():
        for (module, declaration), record in v3.records(corpus).items():
            roots = [n for n in record["nodes"] if n["parent"] is None]
            if len(roots) != 1 or not roots[0]["kind"].endswith("byTactic") or generated(declaration):
                continue
            if (module, declaration) in done:
                continue
            try:
                steps = counter.derive_steps(record["nodes"])
            except ValueError:
                continue
            if any(s["line"] is None for s in steps):
                continue
            source = paths["source"] / (module.replace(".", "/") + ".lean")
            entry = {"corpus": corpus, "module": module, "declaration": declaration,
                     "source": source.relative_to(ROOT).as_posix(), "byLine": roots[0]["line"],
                     "lastLine": max(n["line"] for n in record["nodes"] if n["line"] is not None), "steps": len(steps)}
            out.append(entry)
    out.sort(key=lambda e: (e["corpus"], e["module"], e["byLine"], e["declaration"]))
    return out


def draw_sample() -> list[dict[str, Any]]:
    rng = random.Random(SEED)
    entries = pool()
    short = [e for e in entries if STEPS[0] <= e["steps"] <= STEPS[1]]
    found = {(e["module"], e["declaration"]): constructs_of(e, v3.records(e["corpus"])[(e["module"], e["declaration"])])
             for e in short}
    chosen: list[dict[str, Any]] = []
    taken: set[tuple[str, str]] = set()
    for construct in CONSTRUCTS:
        candidates = [e for e in short if construct in found[(e["module"], e["declaration"])]
                      and (e["module"], e["declaration"]) not in taken]
        for e in rng.sample(candidates, PER_CONSTRUCT):
            chosen.append(e | {"kind": "proof", "construct": construct})
            taken.add((e["module"], e["declaration"]))
    for e in rng.sample([e for e in entries if e["steps"] >= LONG_MIN], WINDOWS):
        steps = counter.derive_steps(v3.records(e["corpus"])[(e["module"], e["declaration"])]["nodes"])
        lines = sorted(s["line"] for s in steps)
        start = rng.randrange(0, len(lines) - WINDOW + 1)
        chosen.append(e | {"kind": "window", "window": [lines[start], lines[start + WINDOW - 1]]})
    out = []
    proofs, windows = 0, 0
    for e in chosen:
        if e["kind"] == "proof":
            proofs += 1
            ident = f"P{proofs:02d}"
        else:
            windows += 1
            ident = f"W{windows}"
        keep = ("corpus", "module", "declaration", "source", "byLine", "lastLine", "kind")
        out.append({"id": ident, **{k: e[k] for k in keep},
                    **({"construct": e["construct"]} if e["kind"] == "proof" else {"window": e["window"]})})
    return out


def source_matches(entry: dict[str, Any]) -> bool:
    """C58: the worksheet's line numbers are the extraction's: the root `by` is on the line the extraction gives."""
    lines = v2.source_lines(entry["source"])
    return 0 < entry["byLine"] <= len(lines) and re.search(r"\bby\b", lines[entry["byLine"] - 1]) is not None


# ---------------------------------------------------------------- the kit


def worksheet_text(sample: list[dict[str, Any]]) -> str:
    out = ["# Worksheet / 题目清单", "",
           "P01 to P12: reconstruct the whole proof. W1 to W3: reconstruct the window's lines; earlier lines are",
           "printed for context. Line numbers are those of the file.", "",
           "P01 至 P12：重建整个证明。W1 至 W3：只重建窗口所含的行；前面的行印出来供参考。行号即文件中的行号。", ""]
    for entry in sample:
        lines = v2.source_lines(entry["source"])
        out += [f"## {entry['id']}. `{entry['declaration']}`", "",
                f"- Module / 模块: `{entry['module']}`", f"- File / 文件: `{entry['source']}`"]
        if entry["kind"] == "proof":
            first, last = max(1, entry["byLine"] - 2), min(len(lines), entry["lastLine"] + 1)
            out += [f"- Reconstruct / 重建: the whole proof, from line {entry['byLine']} / 整个证明，从第 "
                    f"{entry['byLine']} 行开始", ""]
        else:
            low, high = entry["window"]
            first, last = max(entry["byLine"], low - CONTEXT_LINES), min(len(lines), high + 1)
            out += [f"- Reconstruct / 重建: the steps on lines {low} to {high} / 第 {low} 至 {high} 行上的步骤",
                    f"- The proof starts at line {entry['byLine']} / 证明从第 {entry['byLine']} 行开始", ""]
        out += ["```lean"] + [f"{n:5d}  {lines[n - 1]}" for n in range(first, last + 1)] + ["```", ""]
    return "\n".join(out)


def write_answers(sample: list[dict[str, Any]], path: Path = ANSWERS) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    bold, head = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    book = Workbook()
    index = book.active
    index.title = "Index 目录"
    rows = [["Item 题目", "Declaration 声明", "File 文件", "Lines 行"]]
    for e in sample:
        lines = f"{e['byLine']}-{e['lastLine']}" if e["kind"] == "proof" else f"{e['window'][0]}-{e['window'][1]}"
        rows.append([e["id"], e["declaration"], e["source"], lines])
    for r in rows:
        index.append(r)
    for cell in index[1]:
        cell.font, cell.fill = bold, head
    index.append([])
    index.append(["Fill one sheet per item; the rules and two worked examples are in instructions.md."])
    index.append(["每道题填写一个工作表；规则和两道例题见 instructions.md。"])
    for col, width in zip("ABCD", (10, 60, 70, 12)):
        index.column_dimensions[col].width = width
    for e in sample:
        sheet = book.create_sheet(e["id"])
        window = e["kind"] == "window"
        sheet.append([e["id"], e["declaration"]])
        sheet.append(["Lines 行", f"{e['byLine']}-{e['lastLine']}" if not window else f"{e['window'][0]}-{e['window'][1]}"])
        sheet.append([])
        sheet.append(["Steps 步骤"])
        header = ["# 编号", "Line 行号", "Tactic 策略"] + (["Outside? 依赖窗口外 (Y/N)"] if window else []) + ["Notes 备注"]
        sheet.append(header)
        start = sheet.max_row
        for i in range(ROWS):
            sheet.append([i])
        sheet.append([])
        sheet.append(["Edges 边"])
        sheet.append(["From # 起点编号", "To # 终点编号", "Notes 备注"])
        for _ in range(ROWS):
            sheet.append([])
        for r in (1, 4, start, start + ROWS + 2, start + ROWS + 3):
            for cell in sheet[r]:
                cell.font = bold
        for cell in sheet[start]:
            cell.fill = head
        for cell in sheet[start + ROWS + 3]:
            cell.fill = head
        for col, width in zip("ABCDE", (14, 14, 48, 22 if window else 40, 40)):
            sheet.column_dimensions[col].width = width
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    book.save(path)


def read_answers(path: Path, sample: list[dict[str, Any]]) -> dict[str, Any]:
    """The filled workbook as reconstructions: steps in row order, edges as [from, to]."""
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=True)
    proofs = []
    for e in sample:
        sheet = book[e["id"]]
        rows = [[c.value for c in row] for row in sheet.iter_rows()]
        steps_head = next(i for i, r in enumerate(rows) if r and r[0] == "# 编号")
        edges_head = next(i for i, r in enumerate(rows) if r and r[0] == "From # 起点编号")
        window = e["kind"] == "window"
        steps, notes, number = [], [], {}
        for r in rows[steps_head + 1:edges_head - 2]:
            if len(r) < 2 or r[1] in (None, ""):
                continue
            step = {"line": int(r[1]), "tactic": str(r[2] or "")}
            if window:
                step["outside"] = str(r[3] or "").strip().upper().startswith("Y")
            note = r[4] if window else r[3]
            if note:
                notes.append(f"step {r[0]}: {note}")
            number[int(r[0])] = len(steps)  # edges name steps by the # column the person sees
            steps.append(step)
        edges = []
        for r in rows[edges_head + 1:]:
            if len(r) >= 2 and r[0] not in (None, "") and r[1] not in (None, ""):
                edges.append([number[int(r[0])], number[int(r[1])]])
                if len(r) > 2 and r[2]:
                    notes.append(f"edge {r[0]}->{r[1]}: {r[2]}")
        proofs.append({"id": e["id"], "module": e["module"], "declaration": e["declaration"], "steps": steps,
                       "edges": edges, "notes": "; ".join(notes)})
    return {"reconstructor": "a person, by hand", "proofs": proofs}


# ---------------------------------------------------------------- the comparison


def compare_window(entry: dict[str, Any], rebuilt: dict[str, Any]) -> dict[str, Any]:
    """The window's steps (by line and order within a line), the edges between them, and which consume a goal created
    outside the window."""
    low, high = entry["window"]
    steps, parents = v3.derived_graph(entry)
    inside = [i for i in sorted(range(len(steps)), key=lambda i: (steps[i]["line"], i)) if low <= steps[i]["line"] <= high]
    keys = dict(zip(inside, v1.keyed([steps[i]["line"] for i in inside])))
    derived_edges = {(keys[p], keys[c]) for c in inside for p in parents[c] if p in keys}
    derived_outside = {keys[c] for c in inside if any(p not in keys for p in parents[c])}
    rebuilt_keys = v1.keyed([s["line"] for s in rebuilt["steps"]])
    rebuilt_edges = {(rebuilt_keys[i], rebuilt_keys[j]) for i, j in rebuilt["edges"]}
    rebuilt_outside = {rebuilt_keys[i] for i, s in enumerate(rebuilt["steps"]) if s.get("outside")}
    derived_lines = collections.Counter(steps[i]["line"] for i in inside)
    rebuilt_lines = collections.Counter(s["line"] for s in rebuilt["steps"])
    steps_agree = derived_lines == rebuilt_lines
    return {"id": entry["id"], "kind": "window", "corpus": entry["corpus"], "module": entry["module"],
            "declaration": entry["declaration"], "window": entry["window"],
            "derivedSteps": len(inside), "rebuiltSteps": len(rebuilt["steps"]), "stepsAgree": steps_agree,
            "edgesAgree": steps_agree and derived_edges == rebuilt_edges and derived_outside == rebuilt_outside,
            "outsideAgree": derived_outside == rebuilt_outside,
            "onlyDerivedLines": sorted((derived_lines - rebuilt_lines).elements()),
            "onlyRebuiltLines": sorted((rebuilt_lines - derived_lines).elements()),
            "onlyDerivedEdges": sorted([list(a), list(b)] for a, b in derived_edges - rebuilt_edges),
            "onlyRebuiltEdges": sorted([list(a), list(b)] for a, b in rebuilt_edges - derived_edges),
            "onlyDerivedOutside": sorted(list(k) for k in derived_outside - rebuilt_outside),
            "onlyRebuiltOutside": sorted(list(k) for k in rebuilt_outside - derived_outside)}


def compare(reconstructions: dict[str, Any], sample: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_id = {p["id"]: p for p in reconstructions["proofs"]}
    out = []
    for entry in sample:
        rebuilt = by_id.get(entry["id"])
        if rebuilt is None or not rebuilt["steps"]:
            out.append({"id": entry["id"], "kind": entry["kind"], "module": entry["module"],
                        "declaration": entry["declaration"], "missing": True, "stepsAgree": False, "edgesAgree": False})
        elif entry["kind"] == "window":
            out.append(compare_window(entry, rebuilt))
        else:
            row = v3.compare_one(entry | {"stratum": entry["construct"]}, rebuilt)
            out.append({"id": entry["id"], "kind": "proof", "construct": entry["construct"], **row})
    return out


def summarize(results: list[dict[str, Any]], adjudication: dict[str, Any], sample: list[dict[str, Any]]) -> dict[str, Any]:
    verdicts = {a["id"]: a["category"] for a in adjudication.get("verdicts", [])}
    disagreements = [r for r in results if not r["edgesAgree"]]
    categories = collections.Counter(verdicts.get(r["id"]) for r in disagreements)
    unadjudicated = sorted(r["id"] for r in disagreements if r["id"] not in verdicts)
    return {"experiment": "extraction-audit-v0.4", "preregistrationSha256": sha256_file(PREREG),
            "items": len(results), "agree": len(results) - len(disagreements),
            "stepsAgree": sum(1 for r in results if r.get("stepsAgree")),
            "steps": sum(r.get("derivedSteps") or 0 for r in results),
            "proofs": {"items": sum(1 for r in results if r["kind"] == "proof"),
                       "agree": sum(1 for r in results if r["kind"] == "proof" and r["edgesAgree"])},
            "windows": {"items": sum(1 for r in results if r["kind"] == "window"),
                        "agree": sum(1 for r in results if r["kind"] == "window" and r["edgesAgree"])},
            "byConstruct": {c: {"items": sum(1 for r in results if r.get("construct") == c),
                                "agree": sum(1 for r in results if r.get("construct") == c and r["edgesAgree"])}
                            for c in CONSTRUCTS},
            "disagreements": len(disagreements), "categories": {c: categories.get(c, 0) for c in v1.CATEGORIES},
            "unadjudicated": unadjudicated,
            "hypotheses": {"H127": {"supported": None if unadjudicated else categories.get("extractor error", 0) == 0}},
            "checks": {"C58": {"holds": all(source_matches(e) for e in sample)}},
            "reconstructionsSha256": sha256_file(RECONSTRUCTIONS),
            "adjudicationSha256": sha256_file(ADJUDICATION) if ADJUDICATION.exists() else None,
            "resultsSha256": sha256_file(RESULTS)}


# ---------------------------------------------------------------- registration and checks


def registration_payload(sample: list[dict[str, Any]], dev: dict[str, Any]) -> dict[str, Any]:
    return {
        "experiment": "extraction-audit-v0.4",
        "question": "does a person, rebuilding goal-origin graphs by hand from the source and Lean's goal display, agree "
                    "with the corrected step derivation on difficult constructs and on windows of very long proofs",
        "sample": {"rule": f"seed {SEED}; pool: every proof with one root `by` block and a name that is not generated, "
                           "from the library, the first slice or holdout-v0.1's slice, outside extraction-audit-v0.1's, "
                           "v0.2's and v0.3's samples, in (corpus, module, line, declaration) order; for each construct "
                           f"in order, {PER_CONSTRUCT} proofs of {STEPS[0]} to {STEPS[1]} steps drawn uniformly without "
                           "replacement among those with the construct and not yet drawn; then "
                           f"{WINDOWS} proofs of at least {LONG_MIN} steps drawn uniformly, and in each a window of "
                           f"{WINDOW} consecutive steps in line order, its start drawn uniformly, given to the reconstructor "
                           "as the line range from its first step's line to its last's",
                   "constructs": CONSTRUCTS, "count": len(sample),
                   "byKind": dict(collections.Counter(e["kind"] for e in sample)),
                   "extractionSha256": {c: sha256_file(p["extraction"]) for c, p in CORPORA.items()}},
        "protocol": ["the reconstructor is a person, a co-author of the paper, who receives the kit "
                     "(human-kit/README.md, instructions.md, worksheet.md, answers.xlsx): items by location and line "
                     "range only, not their constructs, sizes or shapes",
                     "the reconstructor works from the source and Lean's goal display, may read Lean's and Mathlib's "
                     "source and documentation, opens only the kit's folder and the source files the worksheet names "
                     "(not the repository's extractions, datasets, results, scripts, or this experiment's README, "
                     "sample.json or preregistration.json), runs no extractor or script, and reconstructs without help "
                     "from a language model",
                     "the filled answers.xlsx is converted (--import-hand) and committed as "
                     "reconstructions-by-hand.json before --compare first runs",
                     "--compare matches steps by source line and order within the line, as extraction-audit-v0.1 does, "
                     "and compares edges with the corrected derivation's; in a window it compares the steps on the "
                     "window's lines, the edges between them, and the set of them that consume a goal created by a "
                     "step outside the window (a goal with no origin counts as created by none)",
                     "every disagreement is adjudicated with its evidence (Lean's info tree and goal display) into: "
                     + "; ".join(v1.CATEGORIES)],
        "hypotheses": {"H127": "no disagreement is adjudicated as an extractor error"},
        "checks": {"C58": "every item's root `by` lies on the line the extraction gives, so the worksheet's line "
                          "numbers are the extraction's"},
        "kitSha256": {p.name: sha256_file(p) for p in (GUIDE, INSTRUCTIONS, WORKSHEET, ANSWERS)},
        "implementationSha256": {name: sha256_file(path) for name, path in IMPLEMENTATIONS.items()},
        "developmentChecksBeforeRegistration": dev["notes"],
        "resultsAbsentAtRegistration": True,
        "registeredLocalDate": time.strftime("%Y-%m-%d") + " America/Los_Angeles",
    }


def dev_check() -> dict[str, Any]:
    """The comparison code on earlier, committed material only: extraction-audit-v0.3's reconstructions, and windows
    of its longest proof against the derivation itself and against a perturbed copy."""
    v3_sample = json.loads(v3.SAMPLE.read_text(encoding="utf-8"))
    v3_results = v3.compare(json.loads(v3.RECONSTRUCTIONS.read_text(encoding="utf-8")), v3_sample)
    reproduced = sum(1 for r in v3_results if r["edgesAgree"])
    longest = max(v3_sample, key=lambda e: len(v3.derived_graph(e)[0]))
    steps, parents = v3.derived_graph(longest)
    lines = sorted(s["line"] for s in steps)
    entry = longest | {"id": "dev", "window": [lines[10], lines[10 + WINDOW - 1]]}
    low, high = entry["window"]
    inside = [i for i in sorted(range(len(steps)), key=lambda i: (steps[i]["line"], i)) if low <= steps[i]["line"] <= high]
    position = {i: k for k, i in enumerate(inside)}
    exact = {"steps": [{"line": steps[i]["line"], "outside": any(p not in position for p in parents[i])} for i in inside],
             "edges": [[position[p], position[c]] for c in inside for p in parents[c] if p in position]}
    perturbed = json.loads(json.dumps(exact))
    perturbed["steps"][-1]["outside"] = not perturbed["steps"][-1]["outside"]
    same, changed = compare_window(entry, exact)["edgesAgree"], compare_window(entry, perturbed)["edgesAgree"]
    round_trip = workbook_round_trip(v3_sample[0], entry)
    notes = [f"extraction-audit-v0.3's committed reconstructions through this comparison: {reproduced} of "
             f"{len(v3_results)} agree, as v0.3 reports",
             f"a window of {WINDOW} steps of v0.3's longest proof against the derivation itself agrees ({same}) and "
             f"against a copy with one outside flag changed disagrees ({not changed})",
             "an answer workbook for one v0.3 proof and that window, filled with their derived graphs (one blank row "
             f"left in the proof's steps) and read back, compares as agreeing ({round_trip})",
             "the sample was drawn once to write sample.json and the kit; no derived graph of a sampled item was "
             "displayed"]
    return {"v03Agree": reproduced, "v03Items": len(v3_results), "windowSelf": same, "windowPerturbed": changed,
            "roundTrip": round_trip, "notes": notes}


def workbook_round_trip(proof: dict[str, Any], window: dict[str, Any]) -> bool:
    """write_answers, filled by program with derived graphs, then read_answers and compare: both items must agree."""
    import tempfile
    from openpyxl import load_workbook

    record = v3.records(proof["corpus"])[(proof["module"], proof["declaration"])]
    items = [{"id": "P01", "kind": "proof", "construct": "dev",
              "lastLine": max(n["line"] for n in record["nodes"] if n["line"] is not None),
              **{k: proof[k] for k in ("corpus", "module", "declaration", "source", "byLine")}},
             {"id": "W1", "kind": "window", "window": window["window"], "lastLine": window["window"][1],
              **{k: window[k] for k in ("corpus", "module", "declaration", "source", "byLine")}}]
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "answers.xlsx"
        write_answers(items, path)
        book = load_workbook(path)
        for item in items:
            sheet = book[item["id"]]
            steps, parents = v3.derived_graph(item)
            order = sorted(range(len(steps)), key=lambda i: (steps[i]["line"], i))
            if item["kind"] == "window":
                low, high = item["window"]
                order = [i for i in order if low <= steps[i]["line"] <= high]
            head = next(r for r in range(1, sheet.max_row + 1) if sheet.cell(r, 1).value == "# 编号")
            edges_head = next(r for r in range(1, sheet.max_row + 1) if sheet.cell(r, 1).value == "From # 起点编号")
            gap = 1 if item["kind"] == "proof" else 0  # a blank row after the first step, as a person might leave
            number = {}
            for k, i in enumerate(order):
                row = head + 1 + k + (gap if k > 0 else 0)
                number[i] = sheet.cell(row, 1).value
                sheet.cell(row, 2).value = steps[i]["line"]
                sheet.cell(row, 3).value = steps[i]["kind"]
                if item["kind"] == "window":
                    sheet.cell(row, 4).value = "Y" if any(p not in order for p in parents[i]) else "N"
            r = edges_head + 1
            for c in order:
                for p in parents[c]:
                    if p in number:
                        sheet.cell(r, 1).value, sheet.cell(r, 2).value = number[p], number[c]
                        r += 1
        book.save(path)
        rebuilt = read_answers(path, items)
    results = compare(rebuilt, items)
    return all(row["edgesAgree"] for row in results)


def main() -> int:
    parser = argparse.ArgumentParser()
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--dev", action="store_true")
    mode.add_argument("--register", action="store_true")
    mode.add_argument("--import-hand", type=Path)
    mode.add_argument("--compare", action="store_true")
    mode.add_argument("--check-committed", action="store_true")
    args = parser.parse_args()
    if args.dev:
        print(json.dumps(dev_check(), indent=1))
        return 0
    if args.register:
        if RECONSTRUCTIONS.exists() or RESULTS.exists():
            raise SystemExit("reconstructions or results exist; registration must precede them")
        dev = dev_check()
        if dev["v03Agree"] != dev["v03Items"] or not dev["windowSelf"] or dev["windowPerturbed"]:
            raise SystemExit(f"development checks fail: {dev}")
        sample = draw_sample()
        if not all(source_matches(e) for e in sample):
            raise SystemExit("C58 fails: a worksheet line number differs from the extraction's")
        EXPERIMENT.mkdir(parents=True, exist_ok=True)
        write_lf(SAMPLE, json.dumps(sample, indent=1, ensure_ascii=False) + "\n")
        write_lf(WORKSHEET, worksheet_text(sample) + "\n")
        write_answers(sample)
        write_lf(PREREG, json.dumps(registration_payload(sample, dev), indent=1, sort_keys=True, ensure_ascii=False) + "\n")
        print(f"registered {len(sample)} items: {PREREG}")
        return 0
    prereg = json.loads(PREREG.read_text(encoding="utf-8"))
    for name, path in IMPLEMENTATIONS.items():
        if name != "runner" and prereg["implementationSha256"][name] != sha256_file(path):
            raise SystemExit(f"implementation {name} changed since registration")
    for p in (GUIDE, INSTRUCTIONS, WORKSHEET):
        if prereg["kitSha256"][p.name] != sha256_file(p):
            raise SystemExit(f"the kit's {p.name} changed since registration")
    sample = json.loads(SAMPLE.read_text(encoding="utf-8"))
    if args.import_hand:
        if RECONSTRUCTIONS.exists():
            raise SystemExit("reconstructions-by-hand.json exists; it is committed once")
        write_lf(RECONSTRUCTIONS, json.dumps(read_answers(args.import_hand, sample), indent=1, ensure_ascii=False) + "\n")
        print(f"wrote {RECONSTRUCTIONS}; commit it before --compare")
        return 0
    if not RECONSTRUCTIONS.exists():
        if not all(source_matches(e) for e in sample):
            raise SystemExit("C58 fails")
        print(f"extraction-audit-v0.4: registered, {len(sample)} items, awaiting the reconstructions")
        return 0
    results = compare(json.loads(RECONSTRUCTIONS.read_text(encoding="utf-8")), sample)
    adjudication = json.loads(ADJUDICATION.read_text(encoding="utf-8")) if ADJUDICATION.exists() else {}
    if args.compare:
        write_lf(RESULTS, "".join(json.dumps(r, ensure_ascii=False) + "\n" for r in results))
        summary = summarize(results, adjudication, sample)
        write_lf(SUMMARY, json.dumps(summary, indent=1) + "\n")
        print(json.dumps({k: summary[k] for k in ("items", "agree", "disagreements", "unadjudicated", "hypotheses")}))
        return 0
    if [json.loads(l) for l in RESULTS.read_text(encoding="utf-8").splitlines() if l.strip()] != \
            json.loads(json.dumps(results)):
        raise SystemExit("the committed results do not follow from the committed reconstructions")
    if json.loads(SUMMARY.read_text(encoding="utf-8")) != json.loads(json.dumps(summarize(results, adjudication, sample))):
        raise SystemExit("the committed summary does not follow")
    print(f"extraction-audit-v0.4-check-ok: items={len(results)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
